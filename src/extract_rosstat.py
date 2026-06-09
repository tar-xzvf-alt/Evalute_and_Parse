# coding: utf-8
"""Extract Rosstat regional indicators from .docx and .xlsx files into CSV."""
import os
import pandas as pd
from docx import Document

REGIONS_MAP = {
    'Астраханская область': 'Астраханская область',
    'Астраханская об': 'Астраханская область',
    'Омская область': 'Омская область',
    'Омская об': 'Омская область',
    'Республика Татарстан': 'Татарстан',
    'Респ Татарстан': 'Татарстан',
    'Республика Татарстан (Татарстан)': 'Татарстан',
}

DATA_DIR = 'data/region_stats'

def parse_number(value_str):
    value_str = str(value_str).replace(',', '.').replace(' ', '').replace('\xa0', '').replace('\u2009', '')
    try:
        return float(value_str)
    except ValueError:
        return None

def extract_docx_table(doc, table_indices, year_cols):
    results = {}
    for ti in table_indices:
        if ti >= len(doc.tables):
            continue
        t = doc.tables[ti]
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if not cells:
                continue
            for rn, mapping in REGIONS_MAP.items():
                if rn in cells[0]:
                    values = {}
                    for year, col in year_cols.items():
                        if col < len(cells):
                            val = parse_number(cells[col])
                            if val is not None:
                                values[year] = val
                    if values:
                        results[mapping] = values
                    break
    return results

def extract_xlsx_table(filepath, year_cols):
    """Extract from individual xlsx file (2022 publication format)."""
    results = {}
    df = pd.read_excel(filepath, engine='openpyxl')
    for _, row in df.iterrows():
        region_cell = str(row.iloc[0])
        for rn, mapping in REGIONS_MAP.items():
            if rn in region_cell:
                values = {}
                for year, col in year_cols.items():
                    if col < len(row):
                        val = parse_number(row.iloc[col])
                        if val is not None:
                            values[year] = val
                if values:
                    results[mapping] = values
                break
    return results

def main():
    all_records = []

    # === DOCX publications ===
    # Structure: {publication_year: {indicator: (filename, [table_indices], {year: col})}}
    docx_specs = {
        2021: {
            'unemployment_pct': ('R_03.docx', [36, 49], {2020: 3}),
        },
        2023: {
            'grp_per_capita':            ('R_09.docx',  [3, 4],   {2020: 3, 2021: 4, 2022: 5}),
            'unemployment_pct':          ('R_03.docx',  [36, 37], {2021: 1, 2022: 2}),
            'investment_per_capita':     ('R_10.docx',  [1, 2],   {2020: 3, 2021: 4, 2022: 5}),
            'emissions_thousand_tons':   ('R_08.docx',  [7, 8],   {2020: 3, 2021: 4, 2022: 5}),
            'population':                ('R_02.docx',  [0, 1, 3],{2020: 3, 2021: 4, 2022: 5}),
            'birth_rate':                ('R_02.docx',  [22, 23], {2020: 3, 2021: 4, 2022: 5}),
            'deaths_per_100k':           ('R_02.docx',  [26, 27], {2020: 3, 2021: 4, 2022: 5}),
            'doctors_per_10k':           ('R_06.docx',  [6, 7],   {2020: 9, 2021: 10, 2022: 11}),
            'hospital_beds_per_10k':     ('R_06.docx',  [0, 1],   {2020: 9, 2021: 10, 2022: 11}),
            'captured_pollutants_pct':   ('R_08.docx',  [11, 12], {2020: 3, 2021: 4, 2022: 5}),
            'fresh_water_use':           ('R_08.docx',  [13, 14], {2020: 3, 2021: 4, 2022: 5}),
            'polluted_wastewater':       ('R_08.docx',  [17, 18], {2020: 3, 2021: 4, 2022: 5}),
            'environmental_spending':    ('R_08.docx',  [19, 20], {2020: 3, 2021: 4, 2022: 5}),
            'retail_trade':              ('R_16.docx',  [0, 1],   {2020: 3, 2021: 4, 2022: 5}),
            'average_monthly_wage':      ('R_04-1.docx', [10, 11],{2020: 3, 2021: 4, 2022: 5}),
        },
        2024: {
            'grp_per_capita':            ('R_09.docx',  [3, 4],   {2020: 3, 2021: 4, 2022: 5}),
            'unemployment_pct':          ('R_03.docx',  [36, 37], {2021: 1, 2022: 2, 2023: 3}),
            'investment_per_capita':     ('R_10.docx',  [1, 2],   {2020: 3, 2021: 4, 2022: 5}),
            'emissions_thousand_tons':   ('R_08.docx',  [7, 8],   {2020: 3, 2021: 4, 2022: 5}),
            'population':                ('R_02.docx',  [0, 1, 3],{2020: 3, 2021: 4, 2022: 5, 2023: 6}),
            'birth_rate':                ('R_02.docx',  [22, 23], {2020: 3, 2021: 4, 2022: 5, 2023: 6}),
            'deaths_per_100k':           ('R_02.docx',  [26, 27], {2020: 3, 2021: 4, 2022: 5, 2023: 6}),
            'doctors_per_10k':           ('R_06.docx',  [6, 7],   {2020: 9, 2021: 10, 2022: 11, 2023: 13}),
            'hospital_beds_per_10k':     ('R_06.docx',  [0, 1],   {2020: 9, 2021: 10, 2022: 11, 2023: 13}),
            'captured_pollutants_pct':   ('R_08.docx',  [11, 12], {2020: 3, 2021: 4, 2022: 5, 2023: 6}),
            'fresh_water_use':           ('R_08.docx',  [13, 14], {2020: 3, 2021: 4, 2022: 5, 2023: 6}),
            'polluted_wastewater':       ('R_08.docx',  [17, 18], {2020: 3, 2021: 4, 2022: 5, 2023: 6}),
            'environmental_spending':    ('R_08.docx',  [19, 20], {2020: 3, 2021: 4, 2022: 5, 2023: 6}),
            'retail_trade':              ('R_16.docx',  [0, 1],   {2020: 3, 2021: 4, 2022: 5, 2023: 6}),
            'average_monthly_wage':      ('R_04-1.docx', [10, 11],{2020: 3, 2021: 4, 2022: 5, 2023: 6}),
        },
        2025: {
            'grp_per_capita':            ('R_09.docx',  [3, 4],   {2022: 4, 2023: 5}),
            'unemployment_pct':          ('R_03.docx',  [36, 37], {2023: 2, 2024: 3}),
            'investment_per_capita':     ('R_10.docx',  [1, 2],   {2022: 4, 2023: 5}),
            'emissions_thousand_tons':   ('R_08.docx',  [7, 8],   {2022: 4, 2023: 5}),
            'population':                ('R_02.docx',  [0, 1, 3],{2022: 4, 2023: 5}),
            'birth_rate':                ('R_02.docx',  [22, 23], {2022: 4, 2023: 5}),
            'deaths_per_100k':           ('R_02.docx',  [26, 27], {2022: 4, 2023: 5}),
            'doctors_per_10k':           ('R_06.docx',  [6, 7],   {2022: 9, 2023: 12}),
            'hospital_beds_per_10k':     ('R_06.docx',  [0, 1],   {2022: 9, 2023: 11}),
            'captured_pollutants_pct':   ('R_08.docx',  [11, 12], {2022: 4, 2023: 5}),
            'fresh_water_use':           ('R_08.docx',  [13, 14], {2022: 4, 2023: 5}),
            'polluted_wastewater':       ('R_08.docx',  [17, 18], {2022: 4, 2023: 5}),
            'environmental_spending':    ('R_08.docx',  [19, 20], {2022: 4, 2023: 5}),
            'retail_trade':              ('R_16.docx',  [0, 1],   {2022: 4, 2023: 5}),
            'average_monthly_wage':      ('R_04-1.docx', [10, 11],{2022: 4, 2023: 5}),
        },
    }

    for pub_label, indicators in docx_specs.items():
        pub_path = os.path.join(DATA_DIR, f'Region_Pokaz_{pub_label}')
        if not os.path.isdir(pub_path):
            continue
        for indicator, (fname, table_idx, year_cols) in indicators.items():
            doc_path = os.path.join(pub_path, fname)
            if not os.path.exists(doc_path):
                continue
            doc = Document(doc_path)
            data = extract_docx_table(doc, table_idx, year_cols)
            for region, values in data.items():
                for year, value in values.items():
                    all_records.append({
                        'publication_year': pub_label,
                        'region': region,
                        'year': int(year),
                        'indicator': indicator,
                        'value': value,
                        'unit': '',
                    })

    # === XLSX publication (2022) ===
    xlsx_base = os.path.join(DATA_DIR, 'Region_Pokaz_2022')
    xlsx_specs = {
        'grp_per_capita':            ('09/RR_pokaz_09-02_2022.xlsx', {2020: 5}),
        'investment_per_capita':     ('10/RR_pokaz_10-02_2022.xlsx', {2020: 5, 2021: 6}),
        'emissions_thousand_tons':   ('08/RR_pokaz_08-03_2022.xlsx', {2020: 5, 2021: 6}),
        'population':                ('02/RR_pokaz_02-01_2022.xlsx', {2020: 5, 2021: 6}),
        'doctors_per_10k':           ('06/RR_pokaz_06-03_2022.xlsx', {2020: 11, 2021: 12}),
        'captured_pollutants_pct':   ('08/RR_pokaz_08-05_2022.xlsx', {2020: 5, 2021: 6}),
        'fresh_water_use':           ('08/RR_pokaz_08-06_2022.xlsx', {2020: 5, 2021: 6}),
        'polluted_wastewater':       ('08/RR_pokaz_08-08_2022.xlsx', {2020: 5, 2021: 6}),
        'environmental_spending':    ('08/RR_pokaz_08-09_2022.xlsx', {2020: 4, 2021: 5}),
        'retail_trade':              ('16/RR_pokaz_16-01_2022.xlsx', {2020: 5, 2021: 6}),
    }
    for indicator, (fname, year_cols) in xlsx_specs.items():
        fpath = os.path.join(xlsx_base, fname)
        if not os.path.exists(fpath):
            continue
        data = extract_xlsx_table(fpath, year_cols)
        for region, values in data.items():
            for year, value in values.items():
                all_records.append({
                    'publication_year': 2022,
                    'region': region,
                    'year': int(year),
                    'indicator': indicator,
                    'value': value,
                    'unit': '',
                })

    if not all_records:
        print('Не удалось извлечь данные.')
        return

    df = pd.DataFrame(all_records)
    # Keep latest publication for each (region, year, indicator)
    df = df.sort_values('publication_year').groupby(
        ['region', 'year', 'indicator'], as_index=False
    ).last()

    df_wide = df.pivot_table(
        index=['region', 'year'],
        columns='indicator',
        values='value'
    ).reset_index()
    df_wide = df_wide.sort_values(['region', 'year'])

    output_path = os.path.join(DATA_DIR, 'region_indicators.csv')
    df_wide.to_csv(output_path, index=False, float_format='%.1f')
    print(f'Сохранено: {output_path}')
    print(f'Регионов: {df_wide["region"].nunique()}, лет: {sorted(df_wide["year"].unique())}')
    print(f'Показателей: {len([c for c in df_wide.columns if c not in ("region","year")])}')
    print()
    print(df_wide.to_string(index=False))

if __name__ == '__main__':
    main()
